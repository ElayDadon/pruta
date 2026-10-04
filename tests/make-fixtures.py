#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
tests/make-fixtures.py — מייצר את הקבצים שב-tests/fixtures.

אותו דוח קטן (12 תנועות) נכתב כאן שוב ושוב — כל פעם על ידי תוכנה אחרת ובפורמט אחר —
כדי שהקוראים של פרוטה ייבדקו מול קבצים כמו שבנקים באמת נותנים, ולא רק מול קבצים שפרוטה
עצמה יצרה: xlsxwriter, openpyxl, pandas, LibreOffice (כולל ‎.xls בינארי אמיתי), טקסט בקידודי
Windows, "אקסל" שהוא בעצם HTML, ו-XML Spreadsheet 2003.

הקבצים עצמם כבר נמצאים בתיקייה, והבדיקות (node tests/run.js) לא צריכות את הסקריפט הזה.
הוא כאן כדי שיהיה ברור מאיפה הם באו, וכדי שאפשר יהיה לייצר אותם מחדש:

    pip install xlsxwriter openpyxl pandas      (+ LibreOffice מותקן, בשביל שניים מהקבצים)
    python3 tests/make-fixtures.py tests/fixtures
"""
import csv, datetime as dt, io, os, subprocess, sys, shutil
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(os.path.abspath(__file__)), 'fixtures')
os.makedirs(OUT, exist_ok=True)

# one small, fixed statement: (date, description, amount, balance, reference)
ROWS = [
    (dt.date(2026, 3, 1), 'משכורת חברת הייטק בע"מ', 14250.00, 18320.55, '100234'),
    (dt.date(2026, 3, 2), 'שופרסל דיל רמת גן 123', -312.40, 18008.15, '100235'),
    (dt.date(2026, 3, 2), 'WOLT TEL AVIV', -86.00, 17922.15, '100236'),
    (dt.date(2026, 3, 5), 'העברה לשכר דירה', -4800.00, 13122.15, '100237'),
    (dt.date(2026, 3, 9), 'פז ילו דלק', -251.37, 12870.78, '100238'),
    (dt.date(2026, 3, 12), 'NETFLIX.COM', -54.90, 12815.88, '100239'),
    (dt.date(2026, 3, 15), 'רמי לוי שיווק השקמה', -428.10, 12387.78, '100240'),
    (dt.date(2026, 3, 18), 'חברת החשמל לישראל', -389.00, 11998.78, '100241'),
    (dt.date(2026, 3, 22), 'זיכוי מס הכנסה', 640.00, 12638.78, '100242'),
    (dt.date(2026, 3, 28), 'סופר-פארם, סניף "עזריאלי"', -97.80, 12540.98, '100243'),
    (dt.date(2026, 4, 1), 'משכורת חברת הייטק בע"מ', 14250.00, 26790.98, '100244'),
    (dt.date(2026, 4, 3), 'ארנונה עיריית תל אביב', -612.00, 26178.98, '100245'),
]
META = [('דוח תנועות בחשבון',), ('מספר סניף', '905'), ('שם לקוח', 'ישראל ישראלי'), ('בנק', 'פפר'), ('מספר חשבון', '44718'), ('תקופת הדוח', '01/03/2026 - 03/04/2026')]
HEAD = ['תאריך', 'תיאור', 'סכום', 'מטבע', 'יתרה', 'אסמכתא']
d2 = lambda d: d.strftime('%d/%m/%Y')

# ---- 1. xlsxwriter: shared strings, real date cells, numeric amounts with a currency format
import xlsxwriter
wb = xlsxwriter.Workbook(os.path.join(OUT, 'xlsxwriter-dates.xlsx'))
ws = wb.add_worksheet('תנועות'); ws.right_to_left()
fd = wb.add_format({'num_format': 'dd/mm/yyyy'}); fm = wb.add_format({'num_format': '#,##0.00 [$₪-he-IL];[Red]-#,##0.00 [$₪-he-IL]'}); fb = wb.add_format({'bold': True})
r = 0
for m in META:
    for c, v in enumerate(m): ws.write(r, c, v, fb if c == 0 else None)
    r += 1
r += 1
for c, v in enumerate(HEAD): ws.write(r, c, v, fb)
r += 1
for d, desc, amt, bal, ref in ROWS:
    ws.write_datetime(r, 0, dt.datetime(d.year, d.month, d.day), fd); ws.write(r, 1, desc); ws.write_number(r, 2, amt, fm); ws.write(r, 3, '₪'); ws.write_number(r, 4, bal, fm); ws.write(r, 5, ref)
    r += 1
ws.merge_range(0, 0, 0, 5, META[0][0], fb)
wb.close()

# ---- 2. openpyxl (normal mode): shared strings, datetime cells with time, frozen panes, column widths
import openpyxl
from openpyxl.styles import Font
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'דוח'; ws.sheet_view.rightToLeft = True
for m in META: ws.append(list(m))
ws.append([])
ws.append(HEAD)
for d, desc, amt, bal, ref in ROWS:
    ws.append([dt.datetime(d.year, d.month, d.day, 9, 30), desc, amt, 'ILS', bal, int(ref)])
for row in ws.iter_rows(min_row=9, max_col=1):
    row[0].number_format = 'DD/MM/YYYY HH:MM'
ws.freeze_panes = 'A9'; ws['A1'].font = Font(bold=True); ws.column_dimensions['B'].width = 40
wb.create_sheet('הסבר').append(['הקובץ הופק אוטומטית'])
wb.save(os.path.join(OUT, 'openpyxl-datetime.xlsx'))

# ---- 3. the data on the second sheet, behind a cover sheet; amounts split into debit / credit
wb = openpyxl.Workbook(); cover = wb.active; cover.title = 'שער'
cover.append(['בנק לדוגמה']); cover.append(['הופק בתאריך', dt.datetime(2026, 4, 4)])
ws = wb.create_sheet('עובר ושב')
ws.append(['תאריך ערך', 'תאריך', 'פרטים', 'אסמכתא', 'חובה', 'זכות', 'יתרה בש"ח'])
for d, desc, amt, bal, ref in ROWS:
    ws.append([d + dt.timedelta(days=1), d, desc, ref, -amt if amt < 0 else None, amt if amt > 0 else None, bal])
for row in ws.iter_rows(min_row=2, max_col=2):
    for c in row: c.number_format = 'DD/MM/YYYY'
wb.save(os.path.join(OUT, 'second-sheet-debit-credit.xlsx'))

# ---- 4. the 1904 date system (old Mac Excel)
wb = xlsxwriter.Workbook(os.path.join(OUT, 'date1904.xlsx'), {'date_1904': True})
ws = wb.add_worksheet(); fd = wb.add_format({'num_format': 'dd/mm/yyyy'})
for c, v in enumerate(HEAD): ws.write(0, c, v)
for i, (d, desc, amt, bal, ref) in enumerate(ROWS, 1):
    ws.write_datetime(i, 0, dt.datetime(d.year, d.month, d.day), fd); ws.write(i, 1, desc); ws.write(i, 2, amt); ws.write(i, 3, 'ILS'); ws.write(i, 4, bal); ws.write(i, 5, ref)
wb.close()

# ---- 5. pandas (to_excel)
import pandas as pd
df = pd.DataFrame([{'Date': pd.Timestamp(d), 'Description': desc, 'Amount': amt, 'Balance': bal, 'Reference': ref} for d, desc, amt, bal, ref in ROWS])
df.to_excel(os.path.join(OUT, 'pandas.xlsx'), index=False, sheet_name='Transactions')

# ---- 6. LibreOffice: re-save the xlsxwriter file (a different XML dialect), and write a real binary .xls
def soffice(src, fmt, outdir):
    subprocess.run(['soffice', '--headless', '--norestore', '--convert-to', fmt, '--outdir', outdir, src], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)
tmp = os.path.join(OUT, '_lo'); os.makedirs(tmp, exist_ok=True)
shutil.copy(os.path.join(OUT, 'xlsxwriter-dates.xlsx'), os.path.join(tmp, 'libreoffice.xlsx'))
soffice(os.path.join(tmp, 'libreoffice.xlsx'), 'xls', tmp)                      # → real BIFF
soffice(os.path.join(tmp, 'libreoffice.xls'), 'xlsx', OUT)                      # → xlsx written by LibreOffice
if os.path.exists(os.path.join(tmp, 'libreoffice.xls')): shutil.move(os.path.join(tmp, 'libreoffice.xls'), os.path.join(OUT, 'binary-biff.xls'))
shutil.rmtree(tmp, ignore_errors=True)

# ---- 7. text files in the encodings Hebrew Windows actually produces
def table(sep=','):
    buf = io.StringIO(); w = csv.writer(buf, delimiter=sep, lineterminator='\r\n')
    for m in META: w.writerow(m)
    w.writerow([]); w.writerow(HEAD)
    for d, desc, amt, bal, ref in ROWS: w.writerow([d2(d), desc, '%.2f' % amt, '₪' if sep == ',' else 'ש"ח', '%.2f' % bal, ref])
    return buf.getvalue()
open(os.path.join(OUT, 'windows-1255.csv'), 'wb').write(table().replace('₪', 'ש"ח').encode('cp1255'))
open(os.path.join(OUT, 'utf8-bom.csv'), 'wb').write(b'\xef\xbb\xbf' + table().encode('utf-8'))
open(os.path.join(OUT, 'utf16-tabs.txt'), 'wb').write(b'\xff\xfe' + table('\t').encode('utf-16-le'))
open(os.path.join(OUT, 'utf16-no-bom.csv'), 'wb').write(table().replace('₪', 'ILS').encode('utf-16-le'))
open(os.path.join(OUT, 'windows-1252-latin.csv'), 'wb').write('Date,Description,Amount\r\n05/03/2026,Café Zürich – déjeuner,-12.50\r\n06/03/2026,Señor Taco,-9.90\r\n07/03/2026,Salary,3000.00\r\n08/03/2026,Crème brûlée,-7.20\r\n'.encode('cp1252'))

# ---- 8. "Excel" files that are really HTML — what several Israeli banks hand out with an .xls name
rows_html = ''.join('<tr><td>%s</td><td>%s</td><td dir="ltr">%s</td><td>%s</td><td dir=ltr>%s<td>%s' % (d2(d), desc.replace('"', '&quot;'), ('{:,.2f}'.format(amt)), '&#8362;', ('{:,.2f}'.format(bal)), ref) for d, desc, amt, bal, ref in ROWS)
html = ('<html dir="rtl"><head><meta http-equiv="Content-Type" content="text/html; charset=windows-1255"><style>td{mso-number-format:"\\@"}</style></head><body>'
        '<table><tr><td colspan="6"><b>דוח תנועות בחשבון</b></td></tr><tr><td>מספר חשבון</td><td>44718</td></tr></table><br>'
        '<table border=1><thead><tr>' + ''.join('<th>%s</th>' % h for h in HEAD) + '</tr></thead><tbody>' + rows_html + '</tbody></table>'
        '<script>document.title="x"</script><img src="http://example.invalid/pixel.gif"></body></html>')
open(os.path.join(OUT, 'html-table.xls'), 'wb').write(html.replace('&#8362;', 'ש"ח').encode('cp1255'))
open(os.path.join(OUT, 'html-table-utf8.xls'), 'wb').write(html.replace('charset=windows-1255', 'charset=utf-8').encode('utf-8'))

# ---- 9. XML Spreadsheet 2003
cells = lambda vals: ''.join('<Cell%s><Data ss:Type="%s">%s</Data></Cell>' % (ix, t, v) for ix, t, v in vals)
xml_rows = '<Row>' + cells([('', 'String', h) for h in HEAD]) + '</Row>'
for d, desc, amt, bal, ref in ROWS:
    xml_rows += '<Row>' + cells([('', 'DateTime', d.strftime('%Y-%m-%dT00:00:00.000')), ('', 'String', desc.replace('&', '&amp;').replace('"', '&quot;')), ('', 'Number', repr(amt)), (' ss:Index="5"', 'Number', repr(bal)), ('', 'String', ref)]) + '</Row>'
xml = ('<?xml version="1.0" encoding="UTF-8"?>\n<?mso-application progid="Excel.Sheet"?>\n<Workbook xmlns="urn:schemas-microsoft-com:office:spreadsheet" xmlns:ss="urn:schemas-microsoft-com:office:spreadsheet">'
       '<Worksheet ss:Name="תנועות"><Table>' + xml_rows + '</Table></Worksheet></Workbook>')
open(os.path.join(OUT, 'spreadsheetml-2003.xls'), 'wb').write(xml.encode('utf-8'))

# ---- 10. xlsx whose XML uses a namespace prefix on every element (some .NET exporters)
import zipfile, re
src = zipfile.ZipFile(os.path.join(OUT, 'xlsxwriter-dates.xlsx'))
dst = zipfile.ZipFile(os.path.join(OUT, 'prefixed-namespace.xlsx'), 'w', zipfile.ZIP_DEFLATED)
for item in src.infolist():
    data = src.read(item.filename)
    if item.filename.startswith('xl/') and item.filename.endswith('.xml') and b'spreadsheetml/2006/main' in data:
        t = data.decode('utf-8')
        t = re.sub(r'<(/?)(?![?!])([A-Za-z][\w.-]*)(?=[\s/>])', lambda m: '<%sx:%s' % (m.group(1), m.group(2)), t)
        t = t.replace('xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"', 'xmlns:x="http://schemas.openxmlformats.org/spreadsheetml/2006/main"', 1)
        data = t.encode('utf-8')
    dst.writestr(item, data)
dst.close()
print('\n'.join('%-34s %6d bytes' % (f, os.path.getsize(os.path.join(OUT, f))) for f in sorted(os.listdir(OUT))))
